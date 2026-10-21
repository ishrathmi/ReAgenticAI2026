"""Convert papers and SI files into text chunks the model can read.

Supported: .pdf (text via PyMuPDF, tables via pdfplumber), .xlsx (every sheet as rows),
.docx (paragraphs and tables, no extra dependency). Figures are not interpreted.
"""

from __future__ import annotations

import hashlib
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET

SUPPORTED = {".pdf", ".xlsx", ".docx"}       # file types scanned in local folders (.xml: Europe PMC full text)
PARSER_VERSION = "2"                         # bump when parsing changes, so the paper store re-reads files
SI_NAME_RE = re.compile(r"(_si_|_si\.|suppl|_esm|mmc\d|supporting|supplementary)", re.I)
SECTION_RE = re.compile(
    r"^\s*(?:\d{1,2}(?:\.\d{1,2})*\.?\s+)?(abstract|introduction|background|results and discussion|results|"
    r"discussion|conclusions?|summary|materials and methods|methods and materials|methods|experimental section|"
    r"experimental procedures|experimental|materials|acknowledg(?:e)?ments?|references|notes and references|"
    r"supporting information|supplementary (?:information|materials?|methods))\s*$", re.I | re.M)
CAPTION_RE = re.compile(r"^\s*\[?(captions on this page: )?(Table|TABLE)\s+S?-?[0-9IVX]+", re.M)
DOI_RE = re.compile(r"\b(10\.\d{4,9}/[^\s\"<>,;]+)", re.I)


@dataclass
class Chunk:
    source_file: str
    page: int | None     # 1-based page for PDFs; None for xlsx/docx
    kind: str            # "text" or "table"
    label: str           # human-readable location, e.g. "p4 text 2", "sheet Table S2"
    text: str
    score: int = 0
    section: str = ""            # e.g. "Methods"; set by section labelling or the source's own structure
    chunk_id: int | None = None  # set by the paper store
    paper_id: str = ""           # set by the paper store


@dataclass
class Document:
    doc_id: str
    path: Path
    title: str = ""
    doi: str = ""
    chunks: list = field(default_factory=list)


def parse_pdf(path: str) -> str:
    """Plain text of a PDF (kept for the original scaffold interface)."""
    return "\n\n".join(c.text for c in parse_document(path).chunks if c.kind == "text")


TABLE_ROWS = 4      # data rows per table chunk; set from prefilter.table_rows_per_chunk


def parse_document(path, chunk_chars: int = 2500, table_rows: int | None = None) -> Document:
    global TABLE_ROWS
    if table_rows:
        TABLE_ROWS = int(table_rows)
    path = Path(path)
    doc = Document(doc_id=hashlib.sha1(path.read_bytes()).hexdigest()[:10], path=path)
    suffix = path.suffix.lower()
    if suffix == ".xml":
        from .jats import read_jats
        read_jats(doc, chunk_chars, _split, _split_table)
        _context_for_uncaptioned_tables(doc)
        return doc
    if suffix == ".pdf":
        _read_pdf(doc, chunk_chars)
    elif suffix == ".xlsx":
        _read_xlsx(doc, chunk_chars)
    elif suffix == ".docx":
        _read_docx(doc, chunk_chars)
    else:
        raise ValueError(f"Unsupported file type: {path.name}")
    if not doc.doi:
        head = " ".join(c.text for c in doc.chunks[:3])
        m = DOI_RE.search(head)
        doc.doi = m.group(1).rstrip(".)") if m else ""
    if not doc.title:
        doc.title = path.stem
    _context_for_uncaptioned_tables(doc)
    label_sections(doc)
    return doc


def label_sections(doc: Document) -> None:
    """Give every chunk the section heading it falls under (Abstract, Results, Methods, ...).

    Headings are lines that consist only of a standard section name. A chunk that contains a
    heading is labelled with all sections it spans, e.g. "Results; Methods". Tables take the
    section of the text on their page. SI files without headings are labelled "Supporting information".
    """
    is_si = bool(SI_NAME_RE.search(doc.path.name))
    current = "Supporting information" if is_si else "Front matter"
    page_section = {}
    for c in doc.chunks:
        if c.kind != "text" or c.section:
            continue
        spans = [current]
        for m in SECTION_RE.finditer(c.text):
            current = _section_name(m.group(1))
            spans.append(current)
        c.section = "; ".join(dict.fromkeys(spans))
        if c.page is not None:
            page_section.setdefault(c.page, spans[0])
    for c in doc.chunks:
        if not c.section:
            c.section = page_section.get(c.page, "Supporting information" if is_si else "")


def _section_name(s: str) -> str:
    s = s.strip().lower()
    if "method" in s or s.startswith("experimental") or s == "materials":
        return "Methods"
    if s.startswith("result") or s == "discussion":
        return "Results and discussion" if "discussion" in s else "Results"
    if s.startswith("conclusion") or s == "summary":
        return "Conclusions"
    if "reference" in s:
        return "References"
    if s.startswith("acknowledg"):
        return "Acknowledgements"
    if s.startswith("support") or s.startswith("supplementary"):
        return "Supporting information"
    return s.capitalize()


def _context_for_uncaptioned_tables(doc: Document) -> None:
    """Give tables without a caption the file's other table captions as labelled context.

    Many SI files put several tables under one heading. The model is told the caption was
    not found, so it does not assume the listed caption belongs to this table.
    """
    caps = []
    for c in doc.chunks:
        if c.kind != "text":
            continue
        for m in CAPTION_RE.finditer(c.text):
            line = c.text[m.start():].split("\n", 1)[0][:200].strip(" []")
            if line not in caps:
                caps.append(line)
    if not caps:
        return
    note = "[No caption found for this table. Table captions elsewhere in this file: " + " / ".join(caps[:5]) + "]"
    for c in doc.chunks:
        if c.kind == "table" and not CAPTION_RE.search(c.text):
            c.text = f"{note}\n{c.text}"


# --- helpers -----------------------------------------------------------------

def _clean(text: str) -> str:
    text = text.replace("­", "")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def _split(text: str, max_chars: int) -> list:
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if len(paras) <= 1:                      # PDFs often lack blank lines; fall back to lines
        paras = [p.strip() for p in text.split("\n") if p.strip()]
    chunks, buf = [], ""
    for p in paras:
        while len(p) > max_chars:
            if buf:
                chunks.append(buf)
                buf = ""
            chunks.append(p[:max_chars])
            p = p[max_chars:]
        if buf and len(buf) + len(p) + 1 > max_chars:
            chunks.append(buf)
            buf = ""
        buf = f"{buf}\n{p}" if buf else p
    if buf:
        chunks.append(buf)
    return chunks


def _rows_to_text(rows) -> str:
    out = []
    for row in rows:
        cells = ["" if c is None else re.sub(r"\s+", " ", str(c)).strip() for c in row]
        if any(cells):
            out.append(" | ".join(cells))
    return "\n".join(out)


def _split_table(text: str, max_chars: int) -> list:
    """Split a table into parts of at most TABLE_ROWS data rows (and max_chars), repeating the header.

    Small parts matter: a reply covering 4 rows x 8 columns stays well inside the output
    limit, whereas a whole table often comes back incomplete.
    """
    return _split_rows(text, max_chars, TABLE_ROWS)


def _split_rows(text: str, max_chars: int, max_rows: int) -> list:
    """Split long tables by rows, repeating the header row in every piece."""
    lines = text.split("\n")
    if (len(text) <= max_chars and len(lines) - 1 <= max_rows) or len(lines) < 3:
        return [text]
    header, pieces, buf = lines[0], [], []
    for line in lines[1:]:
        if buf and (len(buf) >= max_rows or len(header) + sum(len(x) + 1 for x in buf) + len(line) > max_chars):
            pieces.append("\n".join([header] + buf))
            buf = []
        buf.append(line)
    if buf:
        pieces.append("\n".join([header] + buf))
    return pieces


def _read_pdf(doc: Document, chunk_chars: int) -> None:
    import fitz  # PyMuPDF

    tables = _pdf_tables(doc.path)            # {page: [(bbox, rows)]}, only tables with real content
    pdf = fitz.open(doc.path)
    doc.title = ((pdf.metadata or {}).get("title") or "").strip()
    captions = {}
    for i, page in enumerate(pdf, start=1):
        full = _clean(page.get_text("text"))
        captions[i] = " / ".join(m.group(0) for m in re.finditer(r"^\s*(Table|TABLE)\s+S?-?[0-9IVX]+.{0,200}", full, re.M))
        boxes = [b for b, _ in tables.get(i, [])]
        # Read the page block by block and leave out blocks inside a table area: the table is
        # sent once, as a table chunk, instead of twice (as page text and as a table).
        blocks = [b for b in page.get_text("blocks") if b[6] == 0 and not _inside(b[:4], boxes)]
        text = _clean("\n".join(b[4].strip() for b in blocks)) if boxes else full
        for j, piece in enumerate(_split(text, chunk_chars), start=1):
            doc.chunks.append(Chunk(doc.path.name, i, "text", f"p{i} text {j}", piece))
    pdf.close()

    for i, page_tables in sorted(tables.items()):
        for j, (_, rows) in enumerate(page_tables, start=1):
            for k, piece in enumerate(_split_table(_rows_to_text(rows), chunk_chars), start=1):
                if piece:
                    label = f"p{i} table {j}" + (f" part {k}" if k > 1 else "")
                    body = f"[captions on this page: {captions[i]}]\n{piece}" if captions.get(i) else piece
                    doc.chunks.append(Chunk(doc.path.name, i, "table", label, body))


def _pdf_tables(path) -> dict:
    """Tables found by pdfplumber, keeping only those with real content (not chemical structures).

    A table counts if it has at least two non-empty rows, at least a quarter of its words are
    numbers (chemical structures misread as tables are mostly atom labels such as 'NH3+ O S'),
    and it covers less than 70% of the page (larger 'tables' are page layouts misread as tables).
    """
    try:
        import pdfplumber
    except ImportError:
        print("  [warn] pdfplumber not installed; PDF tables will be read as plain text only")
        return {}
    found = {}
    try:
        with pdfplumber.open(path) as pl:
            for i, page in enumerate(pl.pages, start=1):
                area = float(page.width * page.height) or 1.0
                for t in page.find_tables():
                    rows = t.extract() or []
                    filled = [r for r in rows if any(c and str(c).strip() for c in r)]
                    words = [w for r in filled for c in r if c for w in str(c).split()]
                    numeric = sum(1 for w in words if re.fullmatch(r"[<>~±]?[-−+]?\d+([.,]\d+)?(%|[eE][-−]?\d+)?", w))
                    x0, top, x1, bottom = t.bbox
                    if (len(filled) >= 2 and words and numeric / len(words) >= 0.25
                            and (x1 - x0) * (bottom - top) / area < 0.7):
                        found.setdefault(i, []).append((t.bbox, rows))
    except Exception as exc:
        print(f"  [warn] table extraction failed for {Path(path).name}: {exc}")
    return found


def _inside(block, boxes, min_overlap: float = 0.5) -> bool:
    """True if at least min_overlap of the text block's area lies inside one of the table boxes."""
    x0, y0, x1, y1 = block
    a = max((x1 - x0) * (y1 - y0), 1e-6)
    for bx0, by0, bx1, by1 in boxes:
        w = min(x1, bx1) - max(x0, bx0)
        h = min(y1, by1) - max(y0, by0)
        if w > 0 and h > 0 and w * h / a >= min_overlap:
            return True
    return False


def _read_xlsx(doc: Document, chunk_chars: int) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(doc.path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        text = _rows_to_text(ws.iter_rows(values_only=True))
        for k, piece in enumerate(_split_table(text, chunk_chars), start=1):
            if piece:
                label = f"sheet '{ws.title}'" + (f" part {k}" if k > 1 else "")
                doc.chunks.append(Chunk(doc.path.name, None, "table", label, piece))
    wb.close()


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def _read_docx(doc: Document, chunk_chars: int) -> None:
    """Walk the document in reading order, including tables inside text boxes.

    Word stores text boxes twice (mc:Choice and mc:Fallback); the Fallback copy is skipped.
    Each table is prefixed with the nearest preceding 'Table N ...' caption.
    """
    with zipfile.ZipFile(doc.path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    paras, tables, used_captions = [], [], set()

    def text_of(el):
        out = []
        for node in el.iter():
            if node.tag == f"{_W}t":
                out.append(node.text or "")
            elif node.tag == f"{_W}tab":
                out.append(" ")
        return "".join(out).strip()

    def walk(el):
        for child in el:
            if child.tag == _MC_FALLBACK:
                continue
            if child.tag == f"{_W}tbl":
                rows = [[text_of(tc) for tc in tr.findall(f"{_W}tc")] for tr in child.findall(f"{_W}tr")]
                caption = next((p for p in reversed(paras[-4:]) if CAPTION_RE.match(p)), "")
                if caption in used_captions:      # a caption belongs to one table only
                    caption = ""
                used_captions.add(caption)
                tables.append((caption, _rows_to_text(rows)))
            elif child.tag == f"{_W}p" and not any(n.tag == f"{_W}tbl" for n in child.iter()):
                t = text_of(child)
                if t:
                    paras.append(t)
            else:
                walk(child)

    walk(root)
    for n, (caption, text) in enumerate(tables, start=1):
        for k, piece in enumerate(_split_table(text, chunk_chars), start=1):
            if piece:
                label = f"table {n}" + (f" part {k}" if k > 1 else "")
                body = f"{caption}\n{piece}" if caption else piece
                doc.chunks.append(Chunk(doc.path.name, None, "table", label, body))
    for j, piece in enumerate(_split("\n\n".join(paras), chunk_chars), start=1):
        doc.chunks.append(Chunk(doc.path.name, None, "text", f"text {j}", piece))
