"""Write extracted records to CSV and a formatted Excel workbook."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

AGENT_COLUMNS = ["evidence_verified", "evidence_match", "unit_in_source", "needs_review", "review_reason",
                 "data_caveats", "model", "extracted_at"]
WIDE = {"evidence": 60, "paper_title": 40, "title": 50, "notes": 40, "sequence": 28,
        "modifications": 25, "solvent": 22, "buffer": 22, "method": 24, "review_reason": 36, "data_caveats": 30,
        "meaning": 90, "source_file": 30}


def columns(schema: dict, records: list) -> list:
    cols = ["record_id"] + list(schema.get("properties", {}))
    cols += [c for c in AGENT_COLUMNS if c not in cols]
    for r in records:                                  # anything else the domain module added
        cols += [k for k in r if k not in cols]
    return cols


def write(records: list, files: list, schema: dict, out_dir: str, basename: str) -> tuple:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    cols = columns(schema, records)
    rows = [{**r, "record_id": i} for i, r in enumerate(records, start=1)]

    csv_path = out / f"{basename}.csv"
    with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:   # utf-8-sig: Excel shows µ, °
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Records"
    _sheet(ws, cols, rows, highlight="needs_review")
    file_cols = list(dict.fromkeys(k for f in files for k in f)) or ["file"]
    _sheet(wb.create_sheet("Files processed"), file_cols, files)
    guide = [{"column": k, "meaning": v.get("description", "")}
             for k, v in schema.get("properties", {}).items()]
    guide += [
        {"column": "evidence_verified", "meaning": "TRUE if the evidence quote was found in the source text (exact or ≥85% fuzzy match)."},
        {"column": "unit_in_source", "meaning": "TRUE if the extracted unit appears in the source passage; FALSE suggests the model supplied a unit the text does not state."},
        {"column": "needs_review", "meaning": "TRUE rows are highlighted: the record is probably wrong. review_reason explains why. Check these first."},
        {"column": "data_caveats", "meaning": "Gaps in the source itself (for example, no unit stated). Not an extraction error; does not set needs_review."},
    ]
    _sheet(wb.create_sheet("Field guide"), ["column", "meaning"], guide)
    xlsx_path = out / f"{basename}.xlsx"
    wb.save(xlsx_path)
    return csv_path, xlsx_path


def _sheet(ws, cols: list, rows: list, highlight: str | None = None) -> None:
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    flag_fill = PatternFill("solid", fgColor="FFF2CC")
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append([_cell(r.get(c)) for c in cols])
        if highlight and r.get(highlight):
            for cell in ws[ws.max_row]:
                cell.fill = flag_fill
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = WIDE.get(c, 14)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2" if "record_id" in cols else "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def _cell(v):
    if isinstance(v, (list, dict)):
        return str(v)
    return v
